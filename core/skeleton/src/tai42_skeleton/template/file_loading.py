"""Content classification + loading for :meth:`ResourceManager.load_file` and the stored-entry reader.

The type of loaded bytes is detected with Magika (a content sniff, more reliable
than an HTTP ``Content-Type`` header, which is often a generic ``octet-stream``),
then matched against a registry keyed by Magika label -> mime substring ->
extension. :func:`classify` is the one rule that decides whether bytes are text, a document
(text extracted from a non-text format) or binary, for every door this module
serves: the stored-entry reader (:func:`classify` on its sample, then
:func:`load_classified` on the whole bytes, with no second detection), and every
load through :func:`load_content` (the resources door, the ``file_loader`` tool). An
empty object is empty text; every type Magika reports as text loads through the
strict UTF-8 decode unless an extracting handler claims it. A media handler passes
the bytes through as a :data:`~tai42_skeleton.template.media.MediaBlock`. Path-based extractors (PDF, DOCX,
XLSX, CSV, PPTX, ODT, EPUB) materialize the bytes to a temp file first.

Backing dependencies are opt-in: ``pip install 'tai42-skeleton[files]'``. Without
them, importing this module raises
:class:`~tai42_skeleton.operations.errors.NotSupportedError` (501, code
``files-not-configured``) whose message is that install line. The server starts
without them: the kinds summary reports the ``files`` row ``off`` with that line, and
each call that loads content refuses.
"""

from __future__ import annotations

import mimetypes
import os
import tempfile
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from functools import lru_cache
from typing import Literal

from tai42_skeleton.operations.errors import NotSupportedError
from tai42_skeleton.template.errors import FILES_NOT_CONFIGURED_CODE, ContentNotLoadableError
from tai42_skeleton.template.media import MediaBlock
from tai42_skeleton.template.settings import file_loading_settings

try:
    import ebooklib
    import pandas as pd
    from bs4 import BeautifulSoup
    from ebooklib import epub
    from fastmcp.utilities.types import Audio, Image
    from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader
    from magika import Magika
    from odf import teletype
    from odf.opendocument import load as odf_load
    from odf.text import P as OdfParagraph
    from pptx import Presentation
    from striprtf.striprtf import rtf_to_text
except ImportError as exc:
    raise NotSupportedError(
        "tai42-skeleton 'files' resource loading requires the 'files' optional "
        "dependencies (magika, pandas, langchain-community, pypdf, docx2txt, "
        "openpyxl, tabulate, beautifulsoup4, python-pptx, striprtf, odfpy, "
        "ebooklib). Install them with: pip install 'tai42-skeleton[files]'",
        extra={"code": FILES_NOT_CONFIGURED_CODE},
    ) from exc

# Magika reads at most this many bytes from each end of its input (the model's
# ``block_size``), so the first and last window of an object classify it as the whole does.
DETECTION_WINDOW_BYTES = 4096


@lru_cache(maxsize=1)
def _magika_client() -> Magika:
    """Return the process-wide Magika client, built once and reused (constructing it loads a model)."""
    return Magika()


@dataclass(frozen=True)
class _Detection:
    """Magika's answer for some bytes: label, mime, extension (with its dot, lowercased) and whether it is text."""

    label: str
    mime: str
    ext: str
    is_text: bool


def _identify(data: bytes, source: str | None) -> _Detection:
    """Inspect ``data`` with Magika; ``source``'s own suffix is the extension hint when Magika cannot classify it."""
    output = _magika_client().identify_bytes(data).output
    label = output.label
    mime = output.mime_type

    ext = ""
    if output.extensions:
        ext = f".{output.extensions[0]}"
    elif label:
        ext = f".{label}"

    if source and (label in ("unknown", "undefined") or not ext):
        url_ext = os.path.splitext(source.split("?")[0])[1]
        if url_ext:
            ext = url_ext

    return _Detection(label=label, mime=mime, ext=ext.lower(), is_text=bool(output.is_text))


def detect(data: bytes, source: str | None) -> tuple[str, str, str]:
    """Inspect ``data`` with Magika and return ``(label, mime, extension)``.

    ``extension`` includes the leading dot (e.g. ``".pdf"``) and is lowercased;
    it may be empty when neither Magika nor ``source`` yields one. When Magika
    cannot classify the content, ``source``'s own suffix is the extension hint.
    """
    detection = _identify(data, source)
    return detection.label, detection.mime, detection.ext


def detect_mime(data: bytes) -> str | None:
    """Return the Magika-detected mime for ``data`` (``None`` when unknown)."""
    return _magika_client().identify_bytes(data).output.mime_type


@contextmanager
def _temp_file(data: bytes, suffix: str) -> Iterator[str]:
    """Materialize ``data`` to a temp file with ``suffix`` for a path-based loader, removing it on exit."""
    fd, path = tempfile.mkstemp(suffix=suffix)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
        yield path
    finally:
        if os.path.exists(path):
            os.remove(path)


def _media_subtype(mime: str | None, source: str | None, kind: str) -> str:
    """Return the media subtype (``png``, ``wav``, ...) for a ``kind`` block from the detected mime.

    ``kind`` is ``image``/``audio``; falls back to a ``source``-suffix guess.
    """
    if mime and mime.startswith(f"{kind}/"):
        return mime.split("/", 1)[1]
    if source:
        guessed, _ = mimetypes.guess_type(source.split("?")[0])
        if guessed and guessed.startswith(f"{kind}/"):
            return guessed.split("/", 1)[1]
    return kind


def _load_pdf(data: bytes, mime: str | None, source: str | None) -> str:
    with _temp_file(data, ".pdf") as path:
        return "\n".join(doc.page_content for doc in PyPDFLoader(path).load())


def _load_docx(data: bytes, mime: str | None, source: str | None) -> str:
    with _temp_file(data, ".docx") as path:
        return "\n".join(doc.page_content for doc in Docx2txtLoader(path).load())


def _load_xlsx(data: bytes, mime: str | None, source: str | None) -> str:
    with _temp_file(data, ".xlsx") as path:
        return pd.read_excel(path).to_markdown(index=False)


def _load_csv(data: bytes, mime: str | None, source: str | None) -> str:
    with _temp_file(data, ".csv") as path:
        return pd.read_csv(path).to_markdown(index=False)


def _load_pptx(data: bytes, mime: str | None, source: str | None) -> str:
    with _temp_file(data, ".pptx") as path:
        lines: list[str] = []
        for slide in Presentation(path).slides:
            lines.extend(
                shape.text_frame.text  # pyright: ignore[reportAttributeAccessIssue]
                for shape in slide.shapes
                if shape.has_text_frame
            )
        return "\n".join(lines)


def _load_odt(data: bytes, mime: str | None, source: str | None) -> str:
    with _temp_file(data, ".odt") as path:
        document = odf_load(path)
        return "\n".join(teletype.extractText(p) for p in document.getElementsByType(OdfParagraph))


def _load_epub(data: bytes, mime: str | None, source: str | None) -> str:
    with _temp_file(data, ".epub") as path:
        book = epub.read_epub(path)
        parts = [
            BeautifulSoup(item.get_content(), "html.parser").get_text()
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT)
        ]
        return "\n".join(parts)


def _load_html(data: bytes, mime: str | None, source: str | None) -> str:
    return BeautifulSoup(data, "html.parser").get_text()


def _decode_utf8(data: bytes, source: str | None) -> str:
    """Strictly decode text bytes: non-UTF-8 bytes raise loudly rather than silently corrupt."""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ContentNotLoadableError(f"{source!r} is detected as text but is not UTF-8: {exc}") from exc


def _load_rtf(data: bytes, mime: str | None, source: str | None) -> str:
    return rtf_to_text(_decode_utf8(data, source))


def _load_text(data: bytes, mime: str | None, source: str | None) -> str:
    return _decode_utf8(data, source)


def _load_image(data: bytes, mime: str | None, source: str | None) -> MediaBlock:
    return Image(data=data, format=_media_subtype(mime, source, "image"))


def _load_audio(data: bytes, mime: str | None, source: str | None) -> MediaBlock:
    return Audio(data=data, format=_media_subtype(mime, source, "audio"))


@dataclass(frozen=True)
class _Handler:
    """One registry entry: the labels, mime substrings and extensions it matches, its loader and its kind.

    The loader turns the bytes into text or a ``MediaBlock``. ``kind`` is ``"text"`` for the
    plain-text decode, ``"extracted"`` for a handler that parses a format into text, and
    ``"media"`` for a pass-through of image or audio bytes.
    """

    labels: tuple[str, ...]
    mime_substrings: tuple[str, ...]
    extensions: tuple[str, ...]
    load: Callable[[bytes, str | None, str | None], str | MediaBlock]
    kind: Literal["text", "extracted", "media"]


_TEXT_HANDLER = _Handler(
    ("txt", "markdown", "json", "xml", "yaml"),
    ("text/", "json", "xml", "yaml"),
    (".txt", ".md", ".json", ".xml", ".yaml", ".yml"),
    _load_text,
    "text",
)

# Order matters for the mime/extension fallback passes: an earlier entry wins.
_REGISTRY: tuple[_Handler, ...] = (
    _Handler(("pdf",), ("pdf",), (".pdf",), _load_pdf, "extracted"),
    _Handler(("docx",), ("wordprocessingml.document",), (".docx", ".doc"), _load_docx, "extracted"),
    _Handler(("xlsx",), ("spreadsheetml.sheet", "ms-excel"), (".xlsx", ".xls"), _load_xlsx, "extracted"),
    _Handler(("csv",), ("csv",), (".csv",), _load_csv, "extracted"),
    _Handler(("pptx",), ("presentationml.presentation", "ms-powerpoint"), (".pptx", ".ppt"), _load_pptx, "extracted"),
    _Handler(("odt",), ("opendocument.text",), (".odt",), _load_odt, "extracted"),
    _Handler(("epub",), ("epub",), (".epub",), _load_epub, "extracted"),
    _Handler(("html", "xhtml"), ("html",), (".html", ".htm"), _load_html, "extracted"),
    _Handler(("rtf",), ("rtf",), (".rtf",), _load_rtf, "extracted"),
    _TEXT_HANDLER,
    _Handler(
        ("png", "jpeg", "gif", "webp"),
        ("image/",),
        (".png", ".jpg", ".jpeg", ".gif", ".webp"),
        _load_image,
        "media",
    ),
    _Handler(("mp3", "wav"), ("audio/",), (".mp3", ".wav"), _load_audio, "media"),
)


def _find_handler(label: str, mime: str | None, ext: str) -> _Handler | None:
    """Resolve a handler by Magika label, then mime substring, then extension; ``None`` when none matches."""
    for handler in _REGISTRY:
        if label in handler.labels:
            return handler
    for handler in _REGISTRY:
        if mime and any(fragment in mime for fragment in handler.mime_substrings):
            return handler
    for handler in _REGISTRY:
        if ext and ext in handler.extensions:
            return handler
    return None


@dataclass(frozen=True)
class Classified:
    """The handler bytes load through (``None``: unsupported), the detection behind it, and the entry kind."""

    handler: _Handler | None
    detection: _Detection
    kind: Literal["text", "document", "binary"]


# Magika's own answer for empty bytes, given without loading the model.
_EMPTY_DETECTION = _Detection(label="empty", mime="inode/x-empty", ext="", is_text=False)


def classify(data: bytes, source: str | None) -> Classified:
    """The one text, document or binary rule for every door this module serves.

    ``data`` is an object's whole bytes, or its first and last
    :data:`DETECTION_WINDOW_BYTES`: Magika reads only those two windows, so both classify
    the object alike and the result loads its whole bytes through :func:`load_classified`.
    Empty bytes are empty text, with no detection. Otherwise a text or extracting
    handler matched by the detection wins; failing that, bytes Magika reports as text
    load through the strict UTF-8 decode; otherwise the media handler matched, or none.
    An extracting handler gives ``"document"`` when the stored bytes are not text
    (PDF, DOCX, XLSX, …) and ``"text"`` when they are (CSV, HTML, RTF).
    """
    if not data:
        return Classified(handler=_TEXT_HANDLER, detection=_EMPTY_DETECTION, kind="text")
    detection = _identify(data, source)
    handler = _find_handler(detection.label, detection.mime, detection.ext)
    if handler is not None and handler.kind == "extracted":
        return Classified(handler=handler, detection=detection, kind="text" if detection.is_text else "document")
    if handler is not None and handler.kind == "text":
        return Classified(handler=handler, detection=detection, kind="text")
    if detection.is_text:
        return Classified(handler=_TEXT_HANDLER, detection=detection, kind="text")
    return Classified(handler=handler, detection=detection, kind="binary")


def check_within_cap(size: int) -> None:
    """Refuse content of ``size`` bytes over ``FILE_LOADING_MAX_BYTES``, before any detection or parse."""
    cap = file_loading_settings().max_bytes
    if size > cap:
        raise ContentNotLoadableError(f"document is {size} bytes, over the {cap}-byte FILE_LOADING_MAX_BYTES cap")


def load_content(data: bytes, source: str | None) -> str | MediaBlock:
    """Classify ``data`` and load it into text or a ``MediaBlock``.

    ``source`` (the originating url/path, when known) supplies the suffix fallback
    for content Magika cannot classify. Raises :class:`ContentNotLoadableError` for
    bytes over the ``FILE_LOADING_MAX_BYTES`` cap (checked BEFORE any detection or
    parse; a storage-id source bypasses ``fetch_url``'s download cap), an unsupported
    type, text bytes that are not UTF-8, and an extracting handler's failure on the
    bytes (an encrypted or truncated PDF, a ragged CSV); the parser's error is the
    cause. A host fault during extraction (``OSError``, ``MemoryError``) propagates.
    """
    check_within_cap(len(data))
    return _load(data, classify(data, source), source)


def load_classified(data: bytes, classified: Classified, source: str | None) -> str | MediaBlock:
    """Load ``data`` through ``classified``, the decision :func:`classify` made on its sample.

    Raises as :func:`load_content` does; the cap is checked before the handler runs.
    """
    check_within_cap(len(data))
    return _load(data, classified, source)


def _load(data: bytes, classified: Classified, source: str | None) -> str | MediaBlock:
    """Run the classified handler on ``data``; an extracting handler's own failure is a ``ContentNotLoadableError``."""
    handler = classified.handler
    detection = classified.detection
    if handler is None:
        raise ContentNotLoadableError(
            f"Unsupported file type: label={detection.label!r} mime={detection.mime!r} extension={detection.ext!r}"
        )
    mime = detection.mime
    if handler.kind != "extracted":
        return handler.load(data, mime, source)
    try:
        return handler.load(data, mime, source)
    except (ContentNotLoadableError, OSError, MemoryError):
        raise
    except Exception as exc:
        detail = str(exc).strip()
        raise ContentNotLoadableError(
            f"{source!r} could not be extracted as {handler.labels[0]}: {type(exc).__name__}"
            + (f": {detail}" if detail else "")
        ) from exc


__all__ = [
    "DETECTION_WINDOW_BYTES",
    "Classified",
    "check_within_cap",
    "classify",
    "detect",
    "detect_mime",
    "load_classified",
    "load_content",
]
